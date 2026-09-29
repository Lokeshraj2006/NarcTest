"""v9: 4 short honest lines on top of the team's v8 (copied from Downloads). Nothing else changed.

Slide 2  beside 'Problem'  : NCRB 2023 NDPS case count (real, sourced)
Slide 2  beside 'Novelty'  : eye check vs lab vs FieldProof comparison
Slide 3  'Not done yet' line: + current stage (lab prototype, TRL 4)
Slide 4  beside 'Cost & Deployment Plan': who pays after the pilot
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE

SRC = 'v8_original.pptx'
OUT = 'K26234_SIH26231_v9.pptx'
p = Presentation(SRC)
E = 914400
GREY = RGBColor(0x55, 0x55, 0x55)


def by_text(slide, text):
    return next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip() == text)


def note(slide, x, y, w, h, lines, size):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap, tf.auto_size = True, MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    for i, line in enumerate(lines):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        for text, bold in line:
            r = par.add_run()
            r.text = text
            r.font.name, r.font.size, r.font.bold = 'Calibri', Pt(size), bold
            r.font.color.rgb = GREY
    return tb


s2, s3, s4 = p.slides[1], p.slides[2], p.slides[3]

# Slide 2 - problem size (NCRB, Crime in India 2023)
h = by_text(s2, 'Problem')
x = h.left / E + 1.50
note(s2, x, h.top / E + 0.10, 8.85 - x, 0.32,
     [[('India: 1,20,010 NDPS cases in 2023 (NCRB) - each one needs a field test.', False)]], 12.5)

# Slide 2 - comparison beside Novelty (inside the Novelty box)
h = by_text(s2, 'Novelty')
x = h.left / E + 1.45
note(s2, x, h.top / E + 0.11, 19.10 - x, 0.32,
     [[('Eye check: guesswork, no record · Lab: accurate but slow · ', False),
       ('FieldProof: on the spot, ₹10, with proof', True)]], 12)

# Slide 3 - honest stage
for shape in s3.shapes:
    if shape.has_text_frame:
        for par in shape.text_frame.paragraphs:
            if par.text.startswith('Not done yet:') and par.runs:
                par.runs[0].text = 'Now: lab prototype (TRL 4). Not done yet: real-kit test with an FSL.'
                for r in par.runs[1:]:
                    r.text = ''

# Slide 4 - who pays after the pilot
h = by_text(s4, 'Cost & Deployment Plan')
x = h.left / E + 3.65
note(s4, x, h.top / E - 0.02, 19.55 - x, 0.52,
     [[('After pilot: ', True), ('paid from the state police / NCB IT budget.', False)],
      [('Open-source, no licence fee, so running cost stays low.', False)]], 13)

p.save(OUT)
print('saved', OUT)
