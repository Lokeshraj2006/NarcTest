"""v18: built on v17. Answers the two 'what to fix' points of the latest review.

Slide 3  real-scan photos + a real PowerPoint table (confusion matrix of 171 checks: 19 real photos, each as taken
         and with 8 computer-made changes, from make_stress_test.py) + plain text results, so a reader or a tool
         that reads slide text can see the numbers (they were only inside a picture before)
Slide 5  map made narrower; next to it the Section 63 certificate page that every PDF report now has
"""
import json
from collections import Counter

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SRC = '../deck_v17/K26234_SIH26231_v17.pptx'
OUT = 'K26234_SIH26231_v18.pptx'
EV = '../ppt_content/evidence/'
HEAD, INK, GREY = RGBColor(0x79, 0x2E, 0x00), RGBColor(0, 0, 0), RGBColor(0x55, 0x55, 0x55)
TAN, OKFILL, WHITE = RGBColor(0xB3, 0x94, 0x60), RGBColor(0xE8, 0xF5, 0xE9), RGBColor(0xFF, 0xFF, 0xFF)
E = 914400
p = Presentation(SRC)


def sid(slide, i):
    return next(s for s in slide.shapes if s.shape_id == i)


def box(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap, tf.auto_size = True, MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    return tf


def lines(tf, rows, size, space=3):
    for i, row in enumerate(rows):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.space_after = Pt(space)
        for text, style in row:
            r = par.add_run()
            r.text = text
            r.font.name, r.font.size, r.font.bold = 'Calibri', Pt(size), style in ('b', 'head')
            r.font.color.rgb = HEAD if style == 'head' else (GREY if style == 'grey' else INK)


s3, s5 = p.slides[2], p.slides[4]

# ------------------------------------------------------------------ Slide 3
data = json.load(open(EV + 'stress_test.json', encoding='utf-8'))
rows = data['rows']
groups = ['Positive card', 'Plain card', 'Trick card']
cols = ['POSITIVE', 'NEGATIVE', 'INCONCLUSIVE']
M = {g: Counter(r['result'] for r in rows if r['group'] == g) for g in groups}
wrong = sum(r['verdict'] == 'wrong' for r in rows)
per_cond = Counter(r['condition'] for r in rows if r['verdict'] == 'wrong')
assert wrong == 0 and not per_cond, 'slide text below says 0 wrong - re-check before using'

old = next(s for s in s3.shapes if s.name == 'Real scans')
old._element.getparent().remove(old._element)
pic = s3.shapes.add_picture(EV + 'real_scans_tiles.png', Inches(10.12), Inches(6.9), width=Inches(3.15))
pic.name = 'Real scan photos'

gf = s3.shapes.add_table(4, 4, Inches(13.42), Inches(6.9), Inches(3.25), Inches(1.2))
gf.name = 'Confusion matrix'
tbl = gf.table
tbl._tbl.tblPr.find(qn('a:tableStyleId')).text = '{5940675A-B579-460E-94D1-54222C63F5DA}'   # plain grid
for c, w in zip(tbl.columns, (1.13, 0.72, 0.72, 0.68)):
    c.width = Inches(w)
head = ['Card \\ App', 'POSITIVE', 'NEGATIVE', 'NOT SURE']
body = [[g] + [str(M[g][c]) for c in cols] for g in groups]
for ri, vals in enumerate([head] + body):
    tbl.rows[ri].height = Inches(0.3)
    for ci, v in enumerate(vals):
        cell = tbl.cell(ri, ci)
        cell.margin_left = cell.margin_right = Inches(0.04)
        cell.margin_top = cell.margin_bottom = Inches(0.01)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        right = ri > 0 and ci == ri
        cell.fill.fore_color.rgb = TAN if ri == 0 else (OKFILL if right else WHITE)
        par = cell.text_frame.paragraphs[0]
        par.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
        r = par.add_run()
        r.text = v
        r.font.name, r.font.size = 'Calibri', Pt(10.5 if ri == 0 else 12)
        r.font.bold = ri == 0 or right
        r.font.color.rgb = INK
cap = box(s3, 13.42, 8.13, 3.25, 0.62)
lines(cap, [[('171 checks: 19 real photos, each as taken and with 8 computer-made changes.', 'grey')]], 11)

res = box(s3, 10.12, 8.9, 6.55, 1.46)
lines(res, [
    [('0 wrong answers in all 171 checks. ', 'b'),
     ('0 wrong in every change: darker, brighter, yellow bulb, cool tube light, tilted, cheap camera, grainy. '
      'Shaky photos all said NOT SURE and asked to retake.', None)],
    [('Colour gap (Delta E), limit 14: ', 'b'),
     ('right colour 4.7 to 8.9, other kit colours 34 to 42, trick cards 16 to 31.', None)],
    [('Real scans: 1 phone (Nothing A015), card on laptop screen, window light, checked by eye.', 'grey')],
], 12.5, space=2)

# ------------------------------------------------------------------ Slide 5
mp = sid(s5, 19)                                     # map: crop to the part with the pin and popup
L, T, R, B = 0.178, 0.10, 0.424, 0.389
bf = mp._element.find('.//' + qn('a:blipFill'))
src = bf.find(qn('a:srcRect'))
for k, v in zip('ltrb', (L, T, R, B)):
    src.set(k, str(int(v * 100000)))
x0 = mp.left
mw = Inches(3.3)
mp.width, mp.height = mw, int(mw * (530 * (1 - T - B)) / (960 * (1 - L - R)))

cx = x0 + mw + Inches(0.2)                           # certificate block to the right of the map
right_edge = sid(s5, 11).left + sid(s5, 11).width    # Future Scope box right edge
hd = box(s5, cx / E, 2.08, (right_edge - cx) / E, 0.42)
lines(hd, [[('Court Certificate - ', 'head'), ('BSA section 63', None)]], 17, space=0)
hd.paragraphs[0].runs[0].font.size = Pt(20.5)
cert = s5.shapes.add_picture(EV + 'section63_page.png', cx, mp.top, height=Inches(1.7))
cert.name = 'Section 63 page'
line = cert.line
line.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
tx = cert.left + cert.width + Inches(0.12)
txt = box(s5, tx / E, mp.top / E, (right_edge - tx) / E, mp.height / E)
lines(txt, [
    [('Done: ', 'b'), ('every PDF report now has this page, filled with photo hash (SHA-256), phone, time and GPS.', None)],
    [('The officer and an expert only sign it.', None)],
    [('Next: ', 'b'), ('written opinion from the FSL and public prosecutor in the pilot.', None)],
], 12, space=3)

p.save(OUT)
print('saved', OUT, '| matrix', {g: dict(M[g]) for g in groups})
