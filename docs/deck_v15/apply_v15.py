"""v15: built on v14. Slide 2 shows the Tamil app for real: the result screenshot in the Prototype Flow row
is now the Tamil result screen (taken on our phone, 29 Sep 2026); caption and UX line mention Tamil."""
from pathlib import Path

from PIL import Image
from pptx import Presentation

SRC = '../deck_v14/K26234_SIH26231_v14.pptx'
OUT = 'K26234_SIH26231_v15.pptx'
SHOT = Path('../ppt_content/evidence/tamil_result_positive.png')
CROP = Path('tamil_result_crop.png')

p = Presentation(SRC)
s2 = p.slides[1]
old = next(s for s in s2.shapes if s.shape_id == 33)            # English result screenshot in the flow row
x, y, w, h = old.left, old.top, old.width, old.height

im = Image.open(SHOT).convert('RGB')                             # 1080 x 2400, portrait
ch = int(im.width * h / w)                                       # same shape as the old picture
top = 70                                                         # skip the phone status bar
im.crop((0, top, im.width, top + ch)).save(CROP)
old._element.getparent().remove(old._element)
pic = s2.shapes.add_picture(str(CROP), x, y, width=w, height=h)
pic.name = 'Tamil result screen'

cap = next(s for s in s2.shapes if s.shape_id == 40)
cap.text_frame.paragraphs[0].runs[0].text = ('Result screen shown in Tamil. Test Log (right) can search and filter '
                                             'by date, officer, kit and result, and download CSV or PDF with QR.')

ux = next(s for s in s2.shapes if s.shape_id == 39).text_frame.paragraphs[2]
ux.runs[-1].text = ' letters, so anyone can read it easily. Works in Tamil too.'

p.save(OUT)
print('saved', OUT)
