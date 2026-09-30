"""Team_189360: the team's 189360 deck layout + content from v16b + fixes from the judge-panel reviews.

Layout, section order and wording style follow 189360_Team_PPT.pdf (the team's version).
Built on the v16b .pptx because it has the same SIH template (logo, oval, footer, fonts).
What changed and why is listed in CHANGES.md next to this file.
Every number comes from docs/ppt_content (validation.md, goals_verified.md,
evidence/stress_test.json, real_tests/records.json) or FieldProof_FULL_PROJECT_NOTES.md.

Run from this folder:   python build_team_189360.py
"""
import copy
import io
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches as I
from pptx.util import Pt

HERE = Path(__file__).resolve().parent
SRC = HERE.parent / 'deck_v16b' / 'K26234_SIH26231_v16b.pptx'
TILES = HERE.parent / 'ppt_content' / 'evidence' / 'real_scans_tiles.png'
TAMPER = HERE.parent / 'ppt_content' / 'evidence' / 'tamper_proof_tall.png'
PHASED = HERE / 'assets' / 'phased_189360.png'      # from 189360_Team_PPT.pdf (icons version)
OUT = HERE / 'Team_189360.pptx'

TITLE, HEAD, INK, GREY = '7A2F00', '974826', '000000', '555555'
TAN, ROWFILL, ROWLINE = 'B2935B', 'FBF6EE', 'D9C4A0'
GREEN, RED, PURPLE, AMBER = '2E7D32', 'C62828', '6A1B9A', '8A5A00'
L, C, R = PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT


# --------------------------------------------------------------------------- helpers
def rgb(h):
    return RGBColor.from_string(h)


def by_id(slide, sid):
    return next(s for s in slide.shapes if s.shape_id == sid)


def kill(slide, *ids):
    for sh in list(slide.shapes):
        if sh.shape_id in ids:
            sh._element.getparent().remove(sh._element)


def runs_of(s):
    """'**Bold** - rest' -> [('Bold', True), (' - rest', False)]"""
    return [(t, i % 2 == 1) for i, t in enumerate(s.split('**')) if t]


def set_bullet(p, ind=0.22):
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', str(int(I(ind))))
    pPr.set('indent', str(-int(I(ind))))
    bf = etree.SubElement(pPr, qn('a:buFont'))
    bf.set('typeface', 'Arial')
    bc = etree.SubElement(pPr, qn('a:buChar'))
    bc.set('char', '•')


def split_paras(paras):
    """str or list of str / (str, opts) -> [(str, opts)], one item per line ('\\n' starts a new paragraph)."""
    out = []
    for para in ([paras] if isinstance(paras, str) else paras):
        text, o = para if isinstance(para, tuple) else (para, {})
        out += [(t, o) for t in text.split('\n')]
    return out


def fill_para(p, s, size=16, color=INK, bold=False, italic=False, align=L, bullet=False,
              space_after=4, link=None):
    p.alignment = align
    if space_after is not None:
        p.space_after = Pt(space_after)
    for t, b in runs_of(s):
        r = p.add_run()
        r.text = t
        f = r.font
        f.name, f.size, f.bold, f.italic = 'Calibri', Pt(size), bold or b, italic
        f.color.rgb = rgb(color)
        if link:
            r.hyperlink.address = link
    if bullet:
        set_bullet(p)


def add_text(slide, x, y, w, h, paras, size=16, color=INK, bold=False, italic=False, align=L,
             anchor=MSO_ANCHOR.TOP, bullet=False, space_after=4, margin=0.03, name=None):
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap, tf.auto_size, tf.vertical_anchor = True, MSO_AUTO_SIZE.NONE, anchor
    tf.margin_left = tf.margin_right = I(margin)
    tf.margin_top = tf.margin_bottom = I(0.02)
    for i, (para, o) in enumerate(split_paras(paras)):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        fill_para(p, para, size=o.get('size', size), color=o.get('color', color), bold=o.get('bold', bold),
                  italic=o.get('italic', italic), align=o.get('align', align),
                  bullet=o.get('bullet', bullet), space_after=o.get('space_after', space_after),
                  link=o.get('link'))
    if name:
        tb.name = name
    return tb


def heading(slide, x, y, w, text, size=25, color=HEAD):
    return add_text(slide, x, y, w, 0.5, text, size=size, color=color, bold=True, space_after=0)


def add_box(slide, x, y, w, h, fill=None, line=TITLE, lw=2.0, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def box_text(s, paras, size=12, color=INK, bold=False, align=C, anchor=MSO_ANCHOR.MIDDLE, margin=0.04):
    tf = s.text_frame
    tf.word_wrap, tf.auto_size, tf.vertical_anchor = True, MSO_AUTO_SIZE.NONE, anchor
    tf.margin_left = tf.margin_right = I(margin)
    tf.margin_top = tf.margin_bottom = I(0.02)
    for i, (para, _) in enumerate(split_paras(paras)):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        fill_para(p, para, size=size, color=color, bold=bold, align=align, space_after=0)
    return s


def set_cell_border(cell, color, w_pt=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        for e in tcPr.findall(qn(tag)):
            tcPr.remove(e)
    for pos, tag in enumerate(('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB')):   # borders must come before the fill
        ln = etree.Element(qn(tag), w=str(int(w_pt * 12700)), cap='flat', cmpd='sng', algn='ctr')
        sf = etree.SubElement(ln, qn('a:solidFill'))
        etree.SubElement(sf, qn('a:srgbClr'), val=color)
        etree.SubElement(ln, qn('a:prstDash'), val='solid')
        tcPr.insert(pos, ln)


def add_table(slide, x, y, col_w, rows, row_h, size=13, head_fill=TAN, border=TITLE, first_col_bold=False,
              head_size=None, name=None):
    n = len(rows)
    row_h = row_h if isinstance(row_h, list) else [row_h] * n
    gf = slide.shapes.add_table(n, len(col_w), I(x), I(y), I(sum(col_w)), I(sum(row_h)))
    tbl = gf.table
    tbl._tbl.tblPr.set('firstRow', '1')
    tbl._tbl.tblPr.set('bandRow', '0')
    for j, w in enumerate(col_w):
        tbl.columns[j].width = I(w)
    for i, h in enumerate(row_h):
        tbl.rows[i].height = I(h)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            o = {}
            if isinstance(val, tuple):
                val, o = val
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(head_fill if i == 0 else o.get('fill', 'FFFFFF'))
            c.margin_left = c.margin_right = I(0.08)
            c.margin_top = c.margin_bottom = I(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.text_frame.word_wrap = True
            for k, t in enumerate(val.split('\n')):
                para = c.text_frame.paragraphs[0] if k == 0 else c.text_frame.add_paragraph()
                fill_para(para, t, size=o.get('size', (head_size or size) if i == 0 else size),
                          color=o.get('color', INK), bold=o.get('bold', i == 0 or (first_col_bold and j == 0)),
                          align=o.get('align', L), space_after=0)
            set_cell_border(c, border)
    if name:
        gf.name = name
    return gf


def line(slide, x1, y1, x2, y2, color=GREY, w=1.0, dash=False):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(w)
    if dash:
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return ln


def down_arrow(slide, cx, y, w=0.34, h=0.26):
    a = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, I(cx - w / 2), I(y), I(w), I(h))
    a.fill.solid()
    a.fill.fore_color.rgb = rgb(TAN)
    a.line.fill.background()
    a.shadow.inherit = False
    return a


def set_runs(shape, paras):
    """Replace paragraph texts in an existing text box, keeping each paragraph's first-run look."""
    ps = shape.text_frame.paragraphs
    for p, new in zip(ps, paras):
        rs = p.runs
        rs[0].text = new
        for r in rs[1:]:
            r._r.getparent().remove(r._r)
    for p in ps[len(paras):]:
        p._p.getparent().remove(p._p)


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def clone_slide(prs, src):
    """Append a copy of `src` (same layout, pictures and links)."""
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    rid = {}
    for k, rel in src.part.rels.items():
        if rel.reltype.endswith(('/slideLayout', '/notesSlide')):
            continue
        rid[k] = (new.part.relate_to(rel.target_ref, rel.reltype, is_external=True) if rel.is_external
                  else new.part.relate_to(rel.target_part, rel.reltype))
    for el in src.shapes._spTree:
        if el.tag in (qn('p:nvGrpSpPr'), qn('p:grpSpPr'), qn('p:extLst')):
            continue
        e = copy.deepcopy(el)
        for node in e.iter():
            for a in (qn('r:embed'), qn('r:id'), qn('r:link')):
                if node.get(a) in rid:
                    node.set(a, rid[node.get(a)])
        new.shapes._spTree.append(e)
    return new


def image_part(slide, shape):
    blip = shape._element.find('.//' + qn('a:blip'))
    return slide.part.related_part(blip.get(qn('r:embed')))


def media(name):
    """A picture from the v16b package, by its media file name."""
    with zipfile.ZipFile(SRC) as z:
        return z.read('ppt/media/' + name)


def picture(slide, blob_or_path, x, y, w=None, h=None, crop=None):
    """crop = (left, top, right, bottom) as fractions of the picture."""
    src = io.BytesIO(blob_or_path) if isinstance(blob_or_path, bytes) else str(blob_or_path)
    im = Image.open(io.BytesIO(blob_or_path) if isinstance(blob_or_path, bytes) else blob_or_path)
    cl, ct, cr, cb = crop or (0, 0, 0, 0)
    aspect = im.width * (1 - cl - cr) / (im.height * (1 - ct - cb))
    if w is None:
        w = h * aspect
    if h is None:
        h = w / aspect
    pic = slide.shapes.add_picture(src, I(x), I(y), I(w), I(h))
    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = cl, ct, cr, cb
    return pic


def move(slide, sid, x=None, y=None, w=None, h=None):
    sh = by_id(slide, sid)
    if x is not None: sh.left = I(x)
    if y is not None: sh.top = I(y)
    if w is not None: sh.width = I(w)
    if h is not None: sh.height = I(h)
    return sh


# --------------------------------------------------------------------------- deck
prs = Presentation(SRC)
s1, s2, s3, s4, s5, s6 = prs.slides
s7 = clone_slide(prs, s6)
s8 = clone_slide(prs, s6)
ORIG = 1.0   # keep the 189360 font sizes where the text fits

set_notes(s1, (
    "OPENING (20 s): We are HexTech, problem statement 26231, Digital Companion for Field Drug Testing. "
    "Field drug tests are read by eye, under any light, with no standard reference and often no record anyone "
    "can check later. That makes the result easy to dispute. FieldProof reads the kit colour with a Rs 10 "
    "colour card, says NOT SURE when it cannot tell, and saves a record that anyone can check with a QR code."))

# ============================================================== slide 2: PROPOSED SOLUTION (189360 layout)
kill(s2, 10, 17, 19, 20, 21, 22, 23, 24, 31, 32, 34, 38, 39, 40, 41, 42, 43)

heading(s2, 1.2, 1.95, 1.6, 'Problem')
add_text(s2, 2.6, 2.08, 6.4, 0.32,
         'India: 1,20,010 NDPS cases registered in 2023 (NCRB, Crime in India 2023)', size=12, color=GREY)
add_text(s2, 1.2, 2.42, 7.85, 2.15, [
    "Officers mainly judge the kit's colour change by looking at it, so lighting and personal judgement "
    'can affect the result.',
    'There is no standard colour value to compare the test result with.',
    'The test does not always leave a clear digital record of who performed it, where, and when, making it '
    'harder to verify later.',
], size=20, bullet=True, space_after=2)
move(s2, 14, x=1.17, y=4.66, w=7.83)
heading(s2, 1.2, 4.76, 3.0, 'Solution')
add_text(s2, 1.2, 5.2, 7.85, 2.65, [
    '**Check the Test** - The app photographs the kit next to a ₹10 colour card and says POSITIVE, NEGATIVE '
    'or NOT SURE.',
    '**Record the Test** - Saves the result with the photo fingerprint, time, GPS and Inspector ID, signed '
    'for a QR check.',
    '**Manage the Records** - The admin dashboard manages Inspectors, test records, NOT SURE reviews and '
    'reports.',
], size=20, bullet=True, space_after=2)
move(s2, 13, x=1.0, y=7.98, w=8.1, h=8.1 / 11.09)          # Test Kit ... Admin Review boxes
for sid, box in {26: (1.2, 9.45, 3.35, 0.43), 27: (4.62, 9.47, 2.3, 0.40), 28: (4.62, 9.42, 2.3, 0.50),
                 25: (7.2, 9.08, 1.16, 1.16), 29: (7.2, 9.08, 1.16, 1.16)}.items():
    move(s2, sid, *box)

add_box(s2, 10.0, 1.87, 7.75, 4.55, line=TITLE, lw=2.25, radius=0.04)
heading(s2, 10.3, 1.98, 1.6, 'Novelty', size=23)
add_text(s2, 11.6, 2.1, 6.0, 0.32, 'Not just a colour reader. A field result anyone can check later.',
         size=13, bold=True, color=GREY, space_after=0)
add_text(s2, 10.3, 2.55, 7.3, 3.8, [
    '**Handles Lighting Variations** - 24 known colours on a ₹10 card correct the light before the kit '
    'colour is read.',
    '**Handles Unclear Results** - Blurry, shiny or in-between photos get NOT SURE with the reason, '
    'never a guess.',
    '**Creates Verifiable Evidence** - Photo fingerprint, GPS, internet time and Inspector ID are signed '
    'and checked by QR.',
    '**Supports Multiple Kits** - New kit colour profiles are added from the Admin website without '
    'changing the app.',
    '**Local Language Support** - Officers can use the app in Tamil.',
], size=19, bullet=True, space_after=6)

move(s2, 30, x=9.38, y=6.55, w=9.9)                        # 'Prototype Flow - real screenshots ...'
ys, hs = 7.12, 3.05
shots = [(media('image8.png'), (0, 0, 0, 0.34)), (media('image9.png'), None),
         (media('image10.png'), (0, 0, 0, 0.18)), (media('image11.png'), (0, 0, 0.38, 0))]
x = 9.45
for k, (blob, crop) in enumerate(shots):
    pic = picture(s2, blob, x, ys, h=hs, crop=crop)
    pic.line.color.rgb = rgb('D0D0D0')
    pic.line.width = Pt(0.75)
    x += pic.width / 914400
    if k < 3:
        move(s2, (35, 36, 37)[k], x=x + 0.035, y=ys + hs / 2 - 0.12, w=0.35, h=0.23)
        x += 0.42

set_notes(s2, (
    "SAY (40 s): Officers judge the colour by eye. A yellow street lamp can make purple look brown. There is "
    "no standard colour and often no record anyone can check later.\n"
    "Our flow: pick the kit, the app runs the kit's read timer, the officer takes one photo of the kit next to "
    "our Rs 10 colour card, and the result comes on screen in about 16 seconds. The record is sealed and gets a QR.\n"
    "Novelty is not colour reading alone. Research has done that (refs 3, 4, 5). What we add is the full path "
    "from field to court: corrected colour, an honest NOT SURE, and a record anyone can check.\n\n"
    "IF ASKED 'Who decides a NOT SURE?': Nobody turns it into POSITIVE. The supervisor can only order a retest "
    "or confirm NOT SURE on the Review page, and the sample goes to FSL as it does today.\n"
    "IF ASKED 'Is this AI?': No AI model, on purpose. Classical colour maths, every result can be explained "
    "with numbers, and it needs no drug photos to train.\n"
    "Source for 1,20,010: NCRB Crime in India 2023 (also quoted in a PIB reply on NDPS cases). It is total NDPS "
    "cases, we do not know how many used a field kit."))

# ============================================================== slide 3: TECHNICAL APPROACH (189360 layout)
kill(s3, 8, 19, 20, 21, 22, 23, 24)

ROWX, ROWW, LBLW = 0.7, 9.1, 1.4
IX0 = ROWX + LBLW + 0.05
IW = ROWX + ROWW - 0.1 - IX0


def arch_row(y, h, label, items, item_h=0.72, gap=0.1):
    add_box(s3, ROWX, y, ROWW, h, fill=ROWFILL, line=ROWLINE, lw=1.0, radius=0.08)
    add_text(s3, ROWX + 0.05, y, LBLW - 0.05, h, label, size=13, bold=True, color=TITLE, align=C,
             anchor=MSO_ANCHOR.MIDDLE, space_after=0)
    w = (IW - gap * (len(items) - 1)) / len(items)
    for k, it in enumerate(items):
        b = add_box(s3, IX0 + k * (w + gap), y + 0.12, w, item_h, fill='FFFFFF', line=TAN, lw=1.0, radius=0.12)
        box_text(b, it, size=12)


y = 2.55
arch_row(y, 0.96, "Officer's phone\n(Flutter app)",
         ['Login and\npick kit', 'Kit read\ntimer', 'Camera with\ncard guide', 'Blur check\non phone',
          'GPS and\nmock flag', 'Photo\nSHA-256'])
down_arrow(s3, IX0 + IW / 2, y + 0.99)
y += 1.3
arch_row(y, 1.46, 'Colour engine\n(server, OpenCV)',
         ['Find card\n(4 corners)', 'Blur and\nglare check', 'Fix colours\n(24 known)', 'CIELAB and\nΔE 2000',
          'Pick closest\nkit colour'])
cw = (IW - 0.2) / 3
for k, (t, col) in enumerate((('POSITIVE', RED), ('NEGATIVE', GREEN), ('NOT SURE + reason', PURPLE))):
    chip = add_box(s3, IX0 + k * (cw + 0.1), y + 0.95, cw, 0.38, fill=col, line=None, radius=0.3)
    box_text(chip, t, size=12, color='FFFFFF', bold=True)
down_arrow(s3, IX0 + IW / 2, y + 1.49)
y += 1.8
arch_row(y, 0.96, 'Sealed record\n(server)',
         ['Internet time\n(NTP, HTTPS)', '22 fields\nsealed', 'SHA-256\nrecord hash', 'Ed25519\nsignature',
          'Chain to last\nrecord'])
down_arrow(s3, IX0 + IW / 2, y + 0.99)
y += 1.3
arch_row(y, 0.96, 'Storage',
         ['PostgreSQL: users, tests,\nkit profiles, reviews', 'Photo store, PDF with QR\nand BSA s.63 page'])
down_arrow(s3, IX0 + IW / 2, y + 0.99)
y += 1.3
arch_row(y, 0.96, 'Admin\nDashboard',
         ['Test log\nand CSV', 'Map', 'Kit editor', 'NOT SURE\nreview', 'QR check\npage'])
add_text(s3, ROWX, y + 1.0, ROWW, 0.3,
         'Everything shown is built. No AI model, on purpose: colour maths that can be explained with numbers.',
         size=12, italic=True, color=GREY, space_after=0)

heading(s3, 10.17, 2.02, 6.4, 'Measured Results')
add_text(s3, 10.2, 2.5, 6.45, 3.62, [
    '**Testing** - 37 computer-made photos in 5 lights and 25 real phone scans.',
    '**Accuracy** - 25/25 computer photos right. Real scans: 14 right, 11 NOT SURE, 0 wrong.',
    '**Colour Check** - ΔE ≤ 14 and a 4-point lead, else NOT SURE. The % shown is a match score, '
    'not a probability.',
    '**Speed** - Server check 2.5 s (5 s max). Whole test about 16 seconds.',
    '**Tamper Test** - All 4/4 changed records were caught.',
], size=17, bullet=True, space_after=8)
add_text(s3, 10.2, 5.72, 6.45, 0.62,
         'Kit colours are printed swatches on a screen, 1 phone. Small test, not proof. Real kits need an FSL '
         'trial (TRL 4).', size=15.5, italic=True, align=C, space_after=0)
line(s3, 10.17, 6.42, 16.55, 6.42, color=TAN, w=1.25)
add_text(s3, 10.17, 6.5, 6.45, 0.34, 'Prototype Validation Using Real Phone Scans (29 Sep)', size=17,
         bold=True, color=TITLE, space_after=0)
picture(s3, TILES, 10.17, 6.95, w=3.6)
add_text(s3, 10.17, 9.25, 3.6, 0.75, '1 phone, 1 room, card on a laptop screen. We checked each card by eye.',
         size=12, color=GREY, space_after=0)
add_text(s3, 13.9, 6.92, 2.75, 0.34, 'All 19 real scans', size=15, bold=True, color=TITLE, space_after=0)
OK_ = {'bold': True, 'color': GREEN, 'fill': 'E8F5E9', 'align': C}
Z = {'color': 'B0B0B0', 'align': C}
add_table(s3, 13.9, 7.3, [1.05, 0.5, 0.5, 0.7], [
    ['Card', ('POS', {'align': C}), ('NEG', {'align': C}), ('NOT SURE', {'align': C})],
    ['Positive (8)', ('5', OK_), ('0', Z), ('3', {'align': C})],
    ['Plain (4)', ('0', Z), ('4', OK_), ('0', Z)],
    ['Trick (7)', ('0', Z), ('0', Z), ('7', OK_)],
], 0.34, size=12, head_fill='F2EBDD', border=ROWLINE)
add_text(s3, 13.9, 8.85, 2.75, 1.4, [
    ('0 wrong answers', {'bold': True, 'color': TITLE, 'size': 13}),
    'All 9 clear scans correct', 'All 7 trick cards said NOT SURE', '3 blurry photos, app asked to retake',
], size=12.5, space_after=1)

set_notes(s3, (
    "SAY (40 s): The card has 24 colours we already know. The app sees how the light changed those colours, "
    "removes that same change from the kit colour, and only then decides. It needs the closest kit colour within "
    "Delta E 14 and at least 4 points ahead of the next one, else it says NOT SURE.\n"
    "Results: 25 real phone scans (6 on 28 Sep, 19 on 29 Sep), 14 right, 11 NOT SURE, 0 wrong. 7 NOT SUREs were "
    "trick cards made to test this, 3 were blurry photos (records 13, 14, 20), 1 was a card-fit reject on 28 Sep.\n"
    "Be honest: printed swatches on a screen, 1 phone. The FSL trial with real kits is next.\n\n"
    "IF ASKED 'What does 76% sure mean?': It is a match score, not a probability. 99 x lead part x closeness "
    "part. Lead part = gap to the next colour / 12, capped at 1. Closeness part = 1 - half of (distance / 14). "
    "We will rename it 'match score' in the app.\n"
    "IF ASKED 'Why Delta E 14?': On real scans the right colour was 4.7 to 8.9 away, trick cards 16 to 31, other "
    "kit colours 34 or more. 14 sits in that gap. The margin to trick cards is thin, so each real kit gets its own "
    "limit from FSL data (backup slide 8).\n"
    "IF ASKED 'Residual error after correction?': That is our card fit number, the average Delta E of the 24 "
    "known colours after correction. 5.9 to 8.7 on real photos of a screen (limit 9), under 1 on computer-made "
    "photos. A paper card test is next.\n"
    "IF ASKED '2.5 s or 16 s?': 2.5 s is the server part (colour check about 0.16 s, most of the rest is the "
    "internet time lookup). 16 s is the whole test by one team member, photo to saved result."))

# ============================================================== slide 4: FEASIBILITY AND VIABILITY (189360 layout)
kill(s4, 10, 14, 21, 54, 55, 56, 57)

add_text(s4, 4.3, 1.98, 5.4, 0.5, [
    '**Legal:** first check only, the FSL lab confirms it.',
    '**Privacy:** no face, fingerprint or suspect name, only a sample code.',
], size=12, color=GREY, space_after=0)
add_text(s4, 0.55, 2.5, 9.1, 3.95, [
    '**Technical** - Full chain works today: app, server, database and dashboard (TRL 4).',
    '**User-Friendly** - One photo, result in about 16 s (tested by our team so far).',
    '**Low Cost** - Existing Android phone, a reusable ₹10 colour card and open-source tools.',
    '**Operational** - Simple photo-based workflow. Officer trial is part of the pilot.',
    '**Scalable** - New kit profiles from the Admin website, after FSL calibration.',
    '**Sustainable** - State police IT cell runs it, FSL keeps kit colours right (proposed).',
], size=19, bullet=True, space_after=10)

picture(s4, PHASED, 0.45, 7.55, w=4.3, crop=(0, 0.146, 0, 0))   # 189360 picture, its own title cropped off
heading(s4, 4.95, 7.08, 4.7, 'Pilot Plan (proposed)', size=20)
add_text(s4, 4.95, 7.5, 4.72, 2.85, [
    '**Where** - 1 district in Tamil Nadu, 20 officers, 3 months, nearest FSL.',
    '**Check 2 things apart** - app vs an expert reading the same kit (target 95%+), and kit vs FSL lab.',
    '**Size** - 60+ tests per kit, so 0 wrong shows error under 5%.',
    '**Pass if** - NOT SURE under 10%, every record passes the QR check.',
    '**Cost** - about ₹9,200, cards and server only (estimate).',
], size=14, bullet=True, space_after=3)

B, P = ('Built', {'bold': True, 'color': GREEN}), ('Planned', {'bold': True, 'color': AMBER})
add_table(s4, 10.0, 2.35, [3.2, 4.45, 1.1], [
    ['Real-world problem', 'What our app does', 'Status'],
    ['Lighting can change the colour seen by the phone.', '24 known colours on the card correct the colour '
                                                          'before analysis.', B],
    ['Blurry or shiny photos can make the result unclear.', 'Detects it, says NOT SURE with the reason and asks '
                                                            'for a retake.', B],
    ['Some colours may fall between two results.', 'Shows NOT SURE instead of guessing. Supervisor can only '
                                                   'order a retest, sample goes to FSL.', B],
    ['Kit colour changes with time.', 'Each kit has a read time, the app timer takes the photo on time.', B],
    ['Photos or records can be edited later.', 'Photo fingerprint, signature and chain catch any change '
                                               '(4/4 in our test).', B],
    ['Someone photographs a screen or a staged colour.', 'In-app camera only today. Screen-photo check and '
                                                         'phone integrity check next.', P],
    ['Different kits and phones show colours differently.', 'Kit colour settings and colour fixed in every '
                                                             'photo. Real kits need FSL calibration.', B],
], [0.5] + [0.6] * 7, size=13, head_size=16, first_col_bold=False)

add_text(s4, 13.35, 7.42, 6.25, 0.5,
         '**After the pilot**, the cost can come from the state police or NCB IT budget. Server about '
         '₹36,000 a year (estimate).', size=12, color=GREY, space_after=0)
for sid, t in {45: 'Existing phone, ₹10 card, open-source software',
               46: 'NIC or state data centre in India (to confirm)',
               47: 'Security updates, FSL keeps kit colours right',
               48: 'Pilot with police and FSL, then scale across locations',
               49: 'New test kits added through the website, no app update'}.items():
    sh = move(s4, sid, y=8.62, h=1.4)
    set_runs(sh, [t])
    sh.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    sh.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

set_notes(s4, (
    "SAY (35 s): It works end to end today as a lab prototype. The pilot we propose is 1 district in Tamil "
    "Nadu, 20 officers, 3 months, with the nearest FSL. We measure 2 things apart: does the app read the kit the "
    "same way a trained expert does, and does the kit agree with the FSL lab. A kit false positive is the kit's "
    "error, not the app's.\n"
    "Every challenge in the table has a status. Built means it is in the code now. Planned means not yet.\n\n"
    "IF ASKED 'Why 60 tests?': Rule of three. With 0 errors in n tests, the true error rate is below 3/n with "
    "95% confidence. 3/60 = 5%. For 1% we would need 300.\n"
    "IF ASKED 'What is not in Rs 9,200?': Kits (already bought today), officers' phones, FSL time, training, "
    "security audit and staff time. We need a quote from the department for those.\n"
    "IF ASKED 'Who owns it after the pilot?': State police IT cell runs the server, FSL owns the kit colour "
    "profiles. MoU is planned, not signed.\n"
    "IF ASKED 'Photo of a screen?': Honest answer, this is open. Our own tests used a screen, which proves "
    "the risk. Plan: screen-photo (moire) check plus Play Integrity. Until then the supervisor sees the photo "
    "and the FSL report is still the final word."))

# ============================================================== slide 5: IMPACT AND BENEFITS (189360 layout)
kill(s5, 11, 22, 28, 31, 32, 33, 34, 35, 36)

set_runs(by_id(s5, 24), ['Same reading in any light', 'Less manual colour checking',
                         'NOT SURE when unclear', 'Easy digital test records',
                         'Quick access to past test records'])
set_runs(by_id(s5, 25), ['Proof of every test: time, place, officer & photo', 'Any change to a record is caught',
                         'PDF with QR and BSA s.63 page', 'Easy search, map and CSV export',
                         'All test data in one system'])

move(s5, 26, x=1.3, w=3.2).text_frame.paragraphs[0].alignment = C
qh = move(s5, 30, x=0.93, y=4.9, w=3.9)
qh.text_frame.paragraphs[0].alignment = L
add_text(s5, 4.85, 4.9, 6.1, 0.42,
         'Judges can scan it on their own phone. For the demo the QR opens a public HTTPS check page '
         '(verify only, nothing else is open).', size=12, color=GREY, space_after=0)
part = image_part(s5, by_id(s5, 10))                  # PDF screenshot: remove the old 127.0.0.1 address
im = Image.open(io.BytesIO(part.blob)).convert('RGB')
ImageDraw.Draw(im).rectangle([400, 251, 645, 264], fill=(255, 255, 255))
buf = io.BytesIO()
im.save(buf, 'PNG')
part._blob = buf.getvalue()

add_text(s5, 12.44, 2.14, 6.6, 0.34, 'Admin can monitor area-wise test results on the map.', size=17,
         space_after=0)

REPORT = 'https://drive.google.com/file/d/1NIR2OyYZRdmEM1R3cz9ghb7rZqyY1D1Z/view?usp=sharing'
heading(s5, 11.62, 4.74, 7.4, 'Evidence Report', size=18)
ev = add_text(s5, 11.62, 5.12, 7.45, 0.66, 'Every PDF report has a filled BSA section 63 page (photo hash, phone, '
              'time, GPS) for the officer and expert to sign. ', size=14, space_after=0)
lr = ev.text_frame.paragraphs[0].add_run()
lr.text = 'Open a full sample report (PDF, 2 pages)'
lr.font.name, lr.font.size, lr.font.underline = 'Calibri', Pt(14), True
lr.font.color.rgb = rgb('0563C1')
lr.hyperlink.address = REPORT
area = add_box(s5, 11.62, 5.12, 7.45, 0.66, fill='FFFFFF', line=None, shape=MSO_SHAPE.RECTANGLE)
area.fill._xPr.find(qn('a:solidFill'))[0].append(etree.Element(qn('a:alpha'), val='0'))
area.click_action.hyperlink.address = REPORT            # click area so the link also works in the PDF
area.name = 'Sample report link area'

add_box(s5, 11.55, 6.0, 7.45, 3.05, line=TITLE, lw=2.25, radius=0.04)
heading(s5, 11.75, 6.08, 7.0, 'Future Scope', size=21.5)
add_text(s5, 11.75, 6.58, 7.1, 2.7, [
    '**Phone Security** - sign each test with the phone’s security chip.',
    '**Lab Comparison** - FSL trial with real kits, compare with lab results.',
    '**Screen-photo Check** - catch photos of a screen, phone integrity check.',
    '**Police Systems** - apply to NIC or MHA for the e-Sakshya and CCTNS link.',
    '**Regional Languages** - Tamil is done, Hindi and voice help next.',
], size=16.5, bullet=True, space_after=9)
add_box(s5, 11.55, 9.22, 7.45, 1.08, line=TITLE, lw=2.25, radius=0.08)
add_text(s5, 11.7, 9.27, 7.2, 0.98,
         '**Offline Mode (Planned, month 1-2)** - Test and photo hash saved and signed on the phone, synced when '
         'the network returns. Duplicates saved once.', size=16, anchor=MSO_ANCHOR.MIDDLE, space_after=0)

set_notes(s5, (
    "SAY (35 s): Every test has time, place, officer and photo. Change even one byte of the photo, or the "
    "result in the database, and the check page says TAMPERED. We tried this on copies of our saved records.\n"
    "Each PDF has a filled BSA 2023 section 63 page. We say it SUPPORTS the certificate. The officer and an "
    "expert still sign, and the FSL report confirms. Never say 'court approved'.\n\n"
    "DEMO (90 s, in this order): 1) scan a positive card in normal light, then under a torch or dim light. "
    "2) scan a trick card, show NOT SURE. 3) show a blurry photo asking for a retake. 4) open the PDF, let a "
    "judge scan the QR on their own phone (run backend/scripts/public_link.py first and re-download the PDF, "
    "the free link lasts 60 min). 5) change 1 byte of a saved copy and show TAMPERED. Keep a screen recording "
    "and a mobile hotspot as backup.\n"
    "IF ASKED 'Offline?': Designed, not built. We will build it together with phone chip signing so an offline "
    "test is signed on the phone before it syncs.\n"
    "IF ASKED 'e-Sakshya?': We have not contacted NIC yet. Our CSV already has time, GPS, officer, result, "
    "photo hash and check link."))

# ============================================================== slide 6: RESEARCH AND REFERENCES
tbl6 = next(s for s in s6.shapes if s.has_table).table
for i, h in enumerate([0.62] + [1.25] * 5):
    tbl6.rows[i].height = I(h)
for i in range(1, 6):
    ref = tbl6.cell(i, 2).text_frame.paragraphs[0]
    ref.line_spacing = 1.05                           # was a fixed 18 pt, which pushed long references out of the row
    for r in ref.runs:
        r.font.size = Pt(12)
    for j in range(3):
        c = tbl6.cell(i, j)
        c.margin_top = c.margin_bottom = I(0.05)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
for (ri, ci, pi), new in {(1, 0, 1): 'Phone camera colours change with light and phone model, and reference colours '
                                     'can help handle this.',
                          (4, 0, 1): 'Reading the colour by eye is subjective, and phone image analysis can give '
                                     'an objective reading.'}.items():
    para = tbl6.cell(ri, ci).text_frame.paragraphs[pi]
    para.runs[0].text = new
    for r in para.runs[1:]:
        r._r.getparent().remove(r._r)
for k in range(5):
    move(s6, 26 + k, y=2.1 + 0.62 + 1.25 * k, h=1.25)
add_text(s6, 0.75, 9.12, 18.5, 0.4,
         '**What we add:** earlier work shows a phone can read test colours [1][3][4][5]. We add the NOT SURE rule, '
         'a signed record anyone can check with a QR, and a BSA section 63 page.', size=14, space_after=0)
add_text(s6, 0.75, 9.66, 18.5, 0.45, 'Scan Smart. Verify Clearly. Protect Evidence.',
         size=20, bold=True, color=TITLE, align=C, space_after=0)
set_notes(s6, (
    "All 5 DOIs open. Reference 1 author: IEEE Xplore and PubMed list the name as 'Karisen', so we cite it as "
    "indexed (the author also publishes as Haakon Karlsen).\n"
    "SAY (20 s, closing): We do not just read a colour, we make the field result checkable. Corrected colour, "
    "honest NOT SURE, and a record anyone can check. Next step is the FSL trial with real kits.\n"
    "IF ASKED 'Why not a handheld Raman device?': Those identify the drug better but are costly and one per "
    "team. FieldProof is for the cheap colour kits officers already carry, and a Raman or lab result can "
    "still be added to the same record later."))

# ============================================================== slides 7 and 8: backup
for s, title, num in ((s7, 'TEST RESULTS IN DETAIL', '7'), (s8, 'SECURITY AND FSL TRIAL', '8')):
    for sh in list(s.shapes):
        if sh.has_table or sh.name.startswith('Research link'):
            sh._element.getparent().remove(sh._element)
    by_id(s, 15).text_frame.paragraphs[0].runs[0].text = title
    by_id(s, 16).text_frame.paragraphs[0].runs[0].text = num
    add_text(s, 4.0, 1.38, 11.4, 0.36, 'Backup slide for judges’ questions', size=14, italic=True,
             color=GREY, align=C, space_after=0)

# ---- slide 7: all tests in one table + colour gap
heading(s7, 0.7, 1.95, 9.6, 'All Tests in One Table', size=23)
add_table(s7, 0.7, 2.45, [2.2, 0.6, 3.75, 3.1], [
    ['Test', 'n', 'Setup', 'Result'],
    ['Computer-made photos', '37', '5 lights: day, bulb, tube light, dim, street lamp. Rules were tuned on these',
     '25/25 good right, 12/12 bad NOT SURE'],
    ['Real scans, 28 Sep', '6', 'Card shown on a phone or laptop screen', '5 right, 1 NOT SURE, 0 wrong'],
    ['Real scans, 29 Sep', '19', '8 positive, 4 plain, 7 trick cards. 1 room, window light',
     '9 right, 10 NOT SURE, 0 wrong'],
    [('All real scans', {'bold': True}), ('25', {'bold': True}), '1 phone (Nothing A015), our team only',
     ('14 right, 11 NOT SURE, 0 wrong', {'bold': True})],
    ['Stress test', '171', '19 real photos x 9 computer-made changes (dark, yellow bulb, tilt, shake, grain)',
     '70 right, 101 NOT SURE, 0 wrong. Sharp photos 69/72 right'],
    ['Tamper test', '5', 'Copies of saved records, 4 changes and 1 untouched', '4/4 changes caught, original passed'],
    ['Speed', '18', 'Server, 6 real photos x 3 runs', '2.5 s typical, 5 s max'],
], [0.42, 0.66, 0.42, 0.66, 0.42, 0.66, 0.42, 0.42], size=13.5)

zb = add_box(s7, 0.7, 7.15, 4.65, 3.1, fill=ROWFILL, line=ROWLINE, lw=1.0)
add_text(s7, 0.85, 7.25, 4.4, 2.95, [
    ('What 0 wrong means', {'bold': True, 'color': TITLE, 'size': 15}),
    '0 wrong in 14 real answers is a small test. By the rule of three the true error rate could still be '
    'up to about 21%.',
    'To show under 5% we need 60+ tests with 0 wrong, so the pilot asks for 60+ tests per kit.',
    'Photo 7 on 28 Sep was rejected when the card fit limit was lower. We quote that live result, not a re-run.',
], size=14.5, space_after=4)
nb = add_box(s7, 5.6, 7.15, 4.7, 3.1, fill=ROWFILL, line=ROWLINE, lw=1.0)
add_text(s7, 5.75, 7.25, 4.45, 2.95, [
    ('What we do not claim', {'bold': True, 'color': TITLE, 'size': 15}),
    ('Accuracy on real reagents (not tested yet)', {'bullet': True}),
    ('Which drug it is (the FSL lab confirms)', {'bullet': True}),
    ('Court approval. It supports the BSA s.63 certificate, the officer and expert still sign', {'bullet': True}),
    ('Results on other phones or with officers (next step)', {'bullet': True}),
], size=15, space_after=3)

# colour gap chart (real scans, Delta E 2000), native shapes
heading(s7, 10.9, 1.95, 8.4, 'Colour Gap on Real Scans (ΔE 2000)', size=23)
X0, X1, VMAX = 13.3, 19.0, 50.0


def vx(v):
    return X0 + (X1 - X0) * v / VMAX


bars = [('Right kit colour', 4.7, 8.9, GREEN, '4.7 to 8.9'),
        ('Trick cards', 16, 31, 'C77C00', '16 to 31'),
        ('Nearest wrong kit colour', 34, 48, RED, '34 to 48')]
for k, (lbl, a, b, col, txt) in enumerate(bars):
    yy = 2.85 + k * 0.55
    add_text(s7, 10.35, yy - 0.02, 2.3, 0.4, lbl, size=13, align=R, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
    add_box(s7, vx(a), yy + 0.04, vx(b) - vx(a), 0.32, fill=col, line=None, shape=MSO_SHAPE.RECTANGLE)
    right = 10 < b < 40                               # label outside the bar, away from the limit line
    lx = vx(b) + 0.08 if right else vx(a) - 1.05
    add_text(s7, lx, yy - 0.02, 1.0, 0.4, txt, size=12, bold=True, color=col, anchor=MSO_ANCHOR.MIDDLE,
             align=L if right else R, space_after=0)
ay = 4.55
line(s7, X0, ay, X1, ay, color=GREY, w=1.0)
for v in range(0, 51, 10):
    line(s7, vx(v), ay, vx(v), ay + 0.08, color=GREY, w=1.0)
    add_text(s7, vx(v) - 0.3, ay + 0.08, 0.6, 0.28, str(v), size=11, color=GREY, align=C, space_after=0)
line(s7, vx(14), 2.78, vx(14), ay, color=TITLE, w=1.75, dash=True)
add_text(s7, vx(14) - 0.6, 2.47, 1.2, 0.3, 'limit 14', size=12, bold=True, color=TITLE, align=C, space_after=0)
add_text(s7, 10.9, 4.98, 8.4, 1.15, [
    'Right colours sat well under 14 and wrong kit colours far above it. The closest trick card was only '
    'about 2 points over the limit, so each real kit gets its own limit from FSL data.',
    '**Card fit** (average error left on the 24 known colours after correction): 5.9 to 8.7 on real photos '
    'of a screen, limit 9. Under 1 on computer-made photos. Paper card test is next.',
], size=12.5, space_after=4)

heading(s7, 10.9, 6.05, 8.4, 'Why the 29 Sep Scans Said NOT SURE', size=20)
add_text(s7, 10.9, 6.47, 8.45, 1.05, [
    ('**3 positive cards** - photo was blurry (records 13, 14, 20), the app asked for a retake.', {'bullet': True}),
    ('**7 trick cards** - 3 matched no kit colour, 4 were stopped earlier by photo checks '
     '(1 blurry, 2 card not found, 1 card fit).', {'bullet': True}),
], size=12.5, space_after=2)
heading(s7, 10.9, 7.12, 8.4, 'Tamper Test on Our Saved Records (29 Sep)', size=20)
picture(s7, TAMPER, 10.9, 7.56, h=2.68)
set_notes(s7, (
    "Use only if asked about numbers. One ledger: 6 scans on 28 Sep + 19 on 29 Sep = 25 real scans. "
    "Right answers 5 + 9 = 14. NOT SURE 1 + 10 = 11. Wrong 0.\n"
    "Round 1 NOT SURE (28 Sep, record 7) was a card-fit reject at 6.16 when the limit was lower. The limit was "
    "later raised to 9 after seeing real photos, so we quote the live result, not a re-run.\n"
    "Round 2 reasons come from the 'as taken' rows of stress_test.json (same engine, same day).\n"
    "The computer-made photos were also used to tune the rules, so 25/25 shows the engine works as designed, "
    "it is not an independent accuracy number.\n"
    "Stress test = computer-changed copies of the 19 real photos, so the 171 checks are not independent tests."))

# ---- slide 8: threats + FSL trial plan
heading(s8, 0.7, 1.95, 10.4, 'Threat Model (what is built, what is open)', size=23)
OPEN = ('Open', {'bold': True, 'color': RED})
FIX = ('Fix planned', {'bold': True, 'color': AMBER})
add_table(s8, 0.7, 2.45, [2.55, 3.3, 3.15, 1.3], [
    ['Threat', 'What stops it today', 'Next step', 'Status'],
    ['Photo or record changed later', 'Photo SHA-256 on phone and server, signed record, chain, QR check',
     'Supervisor counter-seal, officer DSC', ('Built, 4/4 caught', {'bold': True, 'color': GREEN})],
    ['Photo of a screen or staged colour', 'In-app camera only, no gallery upload',
     'Screen-photo (moire) check, Play Integrity', OPEN],
    ['Fake GPS app', 'Android mock-location flag sealed in the record', 'Play Integrity device check',
     ('Partly', {'bold': True, 'color': AMBER})],
    ['Wrong phone clock', 'Server takes internet time, phone time and the gap are also saved',
     'Outside time-stamp (RFC 3161)', ('Built', {'bold': True, 'color': GREEN})],
    ['Insider with the server key', 'Key only on the server, public key is published',
     'Key in a hardware box (HSM), chain head copied outside daily', OPEN],
    ['Server key file lost', 'Server makes a new key, old records then fail the check',
     'Stop and alert instead, keep a key backup', FIX],
    ['Public QR page shows officer name and exact GPS', 'Page is verify only, nothing else is open',
     'Hide the name, round off GPS', FIX],
    ['Test done with no network', 'Not saved today', 'Offline queue built together with phone chip signing',
     ('Planned', {'bold': True, 'color': AMBER})],
    ['One chain for a whole state', 'One chain today (prototype)', 'Chain per device or a daily batch root',
     ('Planned', {'bold': True, 'color': AMBER})],
], [0.42] + [0.66] * 9, size=13.5)

heading(s8, 11.55, 1.95, 7.8, 'FSL Trial to Set Each Kit’s Limit', size=23)
steps = [
    ('Real reagent', 'FSL runs each kit on known samples, 3+ kit lots, photo at the kit’s read time and '
                     'again later to see colour drift.'),
    ('Phones and light', 'Paper card, 3+ phones, daylight, yellow bulb, tube light and street lamp.'),
    ('Measure', 'Card fit on paper, spread of the right colour, distance to the nearest wrong colour.'),
    ('Set limits', 'ΔE limit and lead for each kit from that gap, with a NOT SURE band in between. '
                   'Freeze them before testing.'),
    ('Test fresh samples', 'Samples not used for tuning, compared with the FSL lab. Report agreement '
                           'with a range, not one number.'),
]
for k, (h, t) in enumerate(steps):
    yy = 2.5 + k * 1.2
    dot = add_box(s8, 11.55, yy + 0.05, 0.55, 0.55, fill=TITLE, line=None, shape=MSO_SHAPE.OVAL)
    box_text(dot, str(k + 1), size=16, color='FFFFFF', bold=True)
    add_text(s8, 12.25, yy, 7.1, 1.1, [(h, {'bold': True, 'color': TITLE, 'size': 15, 'space_after': 1}), t],
             size=14, space_after=0)
kb = add_box(s8, 11.55, 8.6, 7.8, 1.65, fill=ROWFILL, line=ROWLINE, lw=1.0)
add_text(s8, 11.7, 8.67, 7.5, 1.55, [
    ('New kit, step by step', {'bold': True, 'color': TITLE, 'size': 15}),
    'FSL calibrates the kit, admin adds the kit profile (it gets a version number), a small field check, then '
    'officers can use it. Kit lot and expiry will be saved in each test record (planned).',
], size=15, space_after=3)
add_text(s8, 0.7, 9.45, 10.3, 0.8,
         '**Legal hooks:** BSA 2023 s.63 (certificate for electronic records, with hash), BNSS 2023 s.105 (recording '
         'of search and seizure), NDPS SSSD Rules 2022, DPDP Act 2023 s.8(5) (keep data safe).',
         size=12.5, color=GREY, space_after=0)
set_notes(s8, (
    "Use only if asked about security or the real-kit plan.\n"
    "Insider question: today whoever holds the server key and the database could rewrite and re-sign the "
    "chain. That is why the key goes into a hardware box (HSM) and the chain head is copied outside every day, "
    "so a rewrite would not match the outside copy.\n"
    "Key loss: today the server makes a new key if the key file is missing, so old records would show TAMPERED. "
    "Fix: stop and alert instead, and keep a backup.\n"
    "Threshold question: our limit 14 came from printed swatches. With real reagents we measure the spread of "
    "the right colour and the nearest wrong colour for each kit, set the limit and lead in that gap, freeze them, "
    "then test on fresh samples against the FSL lab."))

prs.save(OUT)
print('saved', OUT)
