"""v12: built on the team's saved v11 (Downloads). Funding line reworded + font sizes made even.

Slide 2  body text (Problem, Solution, User Experience, Novelty) all 17 pt; grey side notes all 12 pt
Slide 4  Pilot Plan 15 pt like the cost boxes ('Then' line dropped, the phase diagram shows it);
         funding note reworded, 12 pt like the other side notes
"""
from pptx import Presentation
from pptx.util import Inches, Pt

SRC = 'v11_user.pptx'
OUT = 'K26234_SIH26231_v12.pptx'
p = Presentation(SRC)


def sid(slide, i):
    return next(s for s in slide.shapes if s.shape_id == i)


def size(shape, pt, spacing=None):
    for par in shape.text_frame.paragraphs:
        if spacing:
            par.line_spacing = spacing
        for r in par.runs:
            r.font.size = Pt(pt)


s2, s4 = p.slides[1], p.slides[3]

# Slide 2 - one body size, one side-note size
for i in (17, 20, 39):                  # Problem, Solution, User Experience
    size(sid(s2, i), 17)
size(sid(s2, 23), 17, spacing=1.3)      # Novelty: same size, a bit more line gap so the box still looks full
for i in (40, 41, 42):                  # grey side notes
    size(sid(s2, i), 12)

# Slide 4 - pilot plan like the cost boxes, funding note like the other side notes
pilot = sid(s4, 55)
for par in list(pilot.text_frame.paragraphs):
    if par.text.startswith('Then'):
        par._p.getparent().remove(par._p)
size(pilot, 15, spacing=1.1)

note = sid(s4, 56)
p0, p1 = note.text_frame.paragraphs[0], note.text_frame.paragraphs[1]
p0.runs[0].text = 'After the pilot, '
p0.runs[1].text = 'the cost can come from the state police or NCB IT budget.'
p1.runs[0].text = 'No licence fee as it is open source, so running cost stays low.'
for r in p1.runs[1:]:
    r.text = ''
size(note, 12)
note.left, note.width = Inches(13.35), Inches(6.25)

p.save(OUT)
print('saved', OUT)
