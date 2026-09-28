"""v5: answer the second 59/90 review inside the same 6 slides (existing wording kept, blocks added).

Slide 2  + User Experience block (steps, retake reasons, big result text)       -> 'UX section'
Slide 3  + 'Next: real-kit trial with FSL' line under the honest caveat           -> 'real kits'
Slide 4  + Pilot Plan (proposed) next to the phase diagram                        -> 'pilot plan'
Slide 5  + Tamper Test image from a real run (chain, signature, 5 tamper cases)   -> 'signature + verify + log'
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn

SRC = '../deck_v4/K26231_PS_v4.pptx'
OUT = 'K26234_SIH26231_v5.pptx'
TAMPER = '../ppt_content/evidence/tamper_proof.png'
p = Presentation(SRC)
E = 914400


def _walk(shapes):
    for s in shapes:
        yield s
        if s.shape_type == 6:
            yield from _walk(s.shapes)


def sid(slide, i):
    return next(s for s in _walk(slide.shapes) if s.shape_id == i)


def by_text(slide, text):
    return next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip() == text)


def place(shape, x=None, y=None, w=None, h=None):
    if x is not None: shape.left = Inches(x)
    if y is not None: shape.top = Inches(y)
    if w is not None: shape.width = Inches(w)
    if h is not None: shape.height = Inches(h)


def clone(slide, src, x, y, w, h):
    el = copy.deepcopy(src._element)
    el.find('.//' + qn('p:cNvPr')).set('id', str(max(int(e.get('id')) for e in slide._element.iter(qn('p:cNvPr'))) + 1))
    slide.shapes._spTree.insert_element_before(el, 'p:extLst')
    s = next(sh for sh in slide.shapes if sh._element is el)
    place(s, x, y, w, h)
    s.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    s.text_frame.word_wrap = True
    return s


def set_heading(shape, text):
    tf = shape.text_frame
    runs = tf.paragraphs[0].runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for par in tf.paragraphs[1:]:
        par._p.getparent().remove(par._p)
    tf.word_wrap = False


def set_bullets(shape, items, size, spacing=1.1):
    """Rewrite a 'Label - text' bullet box using its first paragraph as the style template."""
    tf = shape.text_frame
    paras = tf.paragraphs
    tpl = next(pa for pa in paras if len(pa.runs) >= 2)._p
    body = tpl.getparent()
    for pa in paras:
        body.remove(pa._p)
    for label, text in items:
        np_ = copy.deepcopy(tpl)
        runs = np_.findall(qn('a:r'))
        for r in runs[2:]:
            np_.remove(r)
        runs[0].find(qn('a:t')).text = label
        runs[1].find(qn('a:t')).text = text
        body.append(np_)
    for pa in tf.paragraphs:
        pa.line_spacing = spacing
        for r in pa.runs:
            r.font.size = Pt(size)


def resize_text(shape, size, spacing):
    for pa in shape.text_frame.paragraphs:
        pa.line_spacing = spacing
        for r in pa.runs:
            r.font.size = Pt(size)


s2, s3, s4, s5 = p.slides[1], p.slides[2], p.slides[3], p.slides[4]

# ------------------------------------------------------------------ Slide 2: User Experience
resize_text(sid(s2, 17), 18, 1.12)               # Problem text
resize_text(sid(s2, 20), 18, 1.12)               # Solution text
place(sid(s2, 19), y=1.98)
place(sid(s2, 17), y=2.45)
place(sid(s2, 14), y=4.22)
place(sid(s2, 21), y=4.32)
place(sid(s2, 20), y=4.38)
ux_head = clone(s2, sid(s2, 21), 1.26, 6.78, 5.0, 0.48)
set_heading(ux_head, 'User Experience')
ux_head.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
ux = clone(s2, sid(s2, 23), 1.08, 7.26, 7.80, 1.40)
set_bullets(ux, [
    ('4 steps', ' - pick kit → photo inside the box → result → sealed. Result ~2.5 s after the photo.'),
    ('Bad photo?', ' - App says NOT SURE with the reason (blur, glare, card not read) and asks to retake.'),
    ('Easy to read', ' - big result word + colour (not colour alone), simple English. Next: timed test with 10 officers.'),
], size=16)

# ------------------------------------------------------------------ Slide 3: real-kit trial is the next step
note = sid(s3, 20).text_frame
nxt = copy.deepcopy(note.paragraphs[0]._p)
nr = nxt.findall(qn('a:r'))
for r in nr[1:]:
    nxt.remove(r)
nr[0].find(qn('a:t')).text = 'Next step (not done yet): real-kit trial with an FSL, in lab + field light.'
note.paragraphs[-1]._p.addnext(nxt)
place(sid(s3, 23), y=7.02)                        # divider
card, ref_head = sid(s3, 13), sid(s3, 22)
proof_head = by_text(s3, 'Proof of Results')
proof = next(s for s in s3.shapes if s.name == 'Proof grid')
pw = 3.10
ph = pw * proof.height / proof.width
place(ref_head, y=7.08)
place(proof_head, x=16.55 - pw, y=7.08, w=pw)
place(proof, x=16.55 - pw, y=7.50, w=pw, h=ph)
place(card, y=7.50 + (ph - card.height / E) / 2)

# ------------------------------------------------------------------ Slide 4: Pilot Plan (proposed)
place(sid(s4, 14), x=0.55, y=7.30, w=4.35, h=4.35 * 3.27 / 5.86)   # phase diagram, smaller
pilot_head = clone(s4, sid(s4, 19), 5.00, 7.10, 4.6, 0.42)
set_heading(pilot_head, 'Pilot Plan (proposed)')
pilot = clone(s4, sid(s4, 21), 4.95, 7.58, 4.72, 2.75)
set_bullets(pilot, [
    ('Where', ' - 1 district in Tamil Nadu, with local police + nearest FSL'),
    ('Who', ' - 20 field officers for 3 months'),
    ('Check', ' - every field result re-tested in the FSL lab'),
    ('Success', ' - 95%+ match with lab, NOT SURE under 10%, every record verifies'),
    ('Then', ' - district → state → national (phases on the left)'),
], size=15.5, spacing=1.15)

# ------------------------------------------------------------------ Slide 5: Tamper Test from a real run
fut, mp = sid(s5, 11), sid(s5, 19)
mh = 2.45
mw = mh * mp.width / mp.height
mx = fut.left / E + (fut.width / E - mw) / 2
dx = Inches(mx) - mp.left
place(mp, x=mx, y=2.62, w=mw, h=mh)
for i in (29, 31):                                # "Map - Geo-Tagged test result" follows the map
    sid(s5, i).left += dx
t_head = clone(s5, sid(s5, 29), fut.left / E, 5.22, 2.0, 0.40)
set_heading(t_head, 'Tamper Test -')
t_sub = clone(s5, sid(s5, 31), fut.left / E + 1.88, 5.26, 5.4, 0.34)
set_heading(t_sub, 'real run on our saved records')
tw = fut.width / E
pic = s5.shapes.add_picture(TAMPER, fut.left, Inches(5.66), width=Inches(tw))
pic.name = 'Tamper proof'
th = pic.height / E
place(fut, y=5.66 + th + 0.14, h=10.25 - (5.66 + th + 0.14))
place(sid(s5, 28), y=5.66 + th + 0.20)            # Future Scope heading
place(sid(s5, 22), y=5.66 + th + 0.62)            # Future Scope text
resize_text(sid(s5, 22), 16, 1.2)

p.save(OUT)
print('saved', OUT)
