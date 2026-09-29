"""v20: built on v19. No 'perfect' claims: '0 wrong' lines reworded to what we saw in a small test, with a real
weakness (NOT SURE on clear cards 38 of 108 in the stress test) and a clear 'real kits can give wrong results' line."""
import json
from collections import Counter

from pptx import Presentation
from pptx.util import Pt

SRC = '../deck_v19/K26234_SIH26231_v19.pptx'
OUT = 'K26234_SIH26231_v20.pptx'
p = Presentation(SRC)

rows = json.load(open('../ppt_content/evidence/stress_test.json', encoding='utf-8'))['rows']
clear = [r for r in rows if r['group'] != 'Trick card']
c = Counter(r['verdict'] for r in rows)
unsure_clear = sum(r['verdict'] == 'not_sure' for r in clear)
print('stress', dict(c), 'clear cards', len(clear), 'not sure on clear', unsure_clear)


def sid(slide, i):
    return next(s for s in slide.shapes if s.shape_id == i)


def set_runs(par, *texts):
    runs = par.runs
    for r, t in zip(runs, texts):
        r.text = t
    for r in runs[len(texts):]:
        r.text = ''


s2, s3 = p.slides[1], p.slides[2]

# Slide 2 - Novelty line promised the app is never wrong
nov = sid(s2, 23).text_frame.paragraphs[1]
for r in nov.runs:
    if 'instead of giving a wrong result' in r.text:
        r.text = r.text.replace('instead of giving a wrong result', 'instead of guessing a result')

# Slide 3 - measured results bullets
goals = sid(s3, 19).text_frame.paragraphs
set_runs(goals[1], 'Accuracy: 25 computer made photos matched. Real scans: 14 of 25 matched, '
                   '11 said NOT SURE (7 were trick cards).')
set_runs(goals[4], 'Tamper Test: we changed 4 saved records on purpose, and the check page caught them.')

# Slide 3 - text under the confusion matrix
res = sid(s3, 30).text_frame.paragraphs
set_runs(res[0], f'Stress test, {len(rows)} checks: ',
         f'{c["right"]} matched the card, {c["not_sure"]} said NOT SURE. No wrong result seen in this small test, '
         f'but clear cards got NOT SURE {unsure_clear} of {len(clear)} times (shaky, cheap camera), '
         f'so retakes are common.')
set_runs(res[2], 'Small test on 1 phone. Real kits can still give wrong results, the FSL trial will check.')
for par in res:
    for r in par.runs:
        r.font.size = Pt(12)

p.save(OUT)
print('saved', OUT)
