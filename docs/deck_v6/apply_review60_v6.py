"""v6: answer the 60/90 review (PPT-only items). Existing wording kept; blocks added.

Slide 2  + caption: the Test Log screenshot is searchable/filterable + CSV/PDF/QR export
Slide 4  + Pilot Plan 'Cost' bullet (per-deployment estimate; per-test cost is already in Feasibility)
Slide 5  QR screenshots made smaller -> Privacy & Retention beside them, Chain of Custody + offline design below
Slide 6  unchanged (references only)
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from lxml import etree

SRC = '../deck_v5/K26234_SIH26231_v5.pptx'
OUT = 'K26234_SIH26231_v6.pptx'
p = Presentation(SRC)
E = 914400
HEAD, GREY, INK = RGBColor(0x79, 0x2E, 0x00), RGBColor(0x6B, 0x6B, 0x6B), RGBColor(0, 0, 0)


def sid(slide, i):
    return next(s for s in slide.shapes if s.shape_id == i)


def place(shape, x=None, y=None, w=None, h=None):
    if x is not None: shape.left = Inches(x)
    if y is not None: shape.top = Inches(y)
    if w is not None: shape.width = Inches(w)
    if h is not None: shape.height = Inches(h)


def box(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap, tf.auto_size = True, MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    return tf


def run(par, text, size, bold=False, color=INK, italic=False):
    r = par.add_run()
    r.text = text
    r.font.name, r.font.size, r.font.bold, r.font.italic = 'Calibri', Pt(size), bold, italic
    r.font.color.rgb = color


def block(slide, x, y, w, h, title, items, size=12, head=15.5):
    tf = box(slide, x, y, w, h)
    par = tf.paragraphs[0]
    run(par, title, head, bold=True, color=HEAD)
    par.space_after = Pt(2)
    for label, text in items:
        par = tf.add_paragraph()
        pPr = par._p.get_or_add_pPr()
        pPr.set('marL', str(int(Inches(0.18))))
        pPr.set('indent', str(-int(Inches(0.15))))
        etree.SubElement(pPr, qn('a:buFont')).set('typeface', 'Arial')
        etree.SubElement(pPr, qn('a:buChar')).set('char', '•')
        par.space_after = Pt(2)
        run(par, label, size, bold=True)
        run(par, text, size)


s2, s4, s5 = p.slides[1], p.slides[3], p.slides[4]

# ------------------------------------------------------------------ Slide 2: searchable log caption
tf = box(s2, 9.40, 6.17, 9.85, 0.24)
run(tf.paragraphs[0], 'Test Log (right side): search and filter by date, officer, kit, result. Download as CSV, PDF with QR.',
    11.5, color=GREY, italic=True)

# ------------------------------------------------------------------ Slide 4: pilot cost (per deployment)
pilot = next(s for s in s4.shapes if s.has_text_frame and s.text_frame.text.startswith('Where'))
last = pilot.text_frame.paragraphs[-1]._p
cost = copy.deepcopy(last)
rs = cost.findall(qn('a:r'))
rs[0].find(qn('a:t')).text = 'Cost'
rs[1].find(qn('a:t')).text = ' - around ₹10,000 for 3 months: cards ₹200 + server ~₹3,000/month (rough estimate)'
last.addnext(cost)
for par in pilot.text_frame.paragraphs:
    for r in par.runs:
        r.font.size = Pt(14)
pilot.height = Inches(10.30 - pilot.top / E)

# ------------------------------------------------------------------ Slide 5: smaller QR screenshots + 2 blocks
pdf, arrow, ver = sid(s5, 10), sid(s5, 15), sid(s5, 16)
S, y0, x0 = 3.40, 5.40, 0.78
k = S / (pdf.height / E)
ax, ay = x0 + (arrow.left / E - pdf.left / E) * k, y0 + (arrow.top / E - pdf.top / E) * k
place(arrow, ax, ay, arrow.width / E * k, arrow.height / E * k)
place(pdf, x0, y0, pdf.width / E * k, S)
vw = ver.width / E * S / (ver.height / E)
place(ver, ax + arrow.width / E + 0.10, y0, vw, S)
right = ver.left / E + vw
qh = sid(s5, 30)                                     # "QR-Based Record Verification" centred over the pair
place(qh, x=(x0 + right) / 2 - qh.width / E / 2)

block(s5, right + 0.28, 4.95, 10.95 - right - 0.28, 3.95, 'Privacy & How Long We Keep Data', [
    ('Data', ' - no fingerprint or face data, no suspect name (only a sample code). Officer ID, GPS, time kept as proof.'),
    ('Who sees', ' - officer: only own tests; supervisor/admin: all tests; public QR: only the check page (hiding officer name: planned).'),
    ('Law', ' - police work does not need consent (DPDP s.7(c), s.17(1)(c)), but data must be kept safe (s.8(5)).'),
    ('Keep', ' - till the case and appeal are over, then photo is deleted (planned).'),
])
block(s5, x0, y0 + S + 0.08, 10.95 - x0, 10.38 - (y0 + S + 0.08), 'Chain of Custody (safe for court)', [
    ('Key', ' - secret key only on our server, never in the app; anyone can check using the public key. '
            "Next: hardware key box (HSM) + officer's own DSC."),
    ('Check', ' - QR page checks the photo and record again, plus the signature and chain link (see Tamper Test).'),
    ('Court', ' - gives the SHA-256 hash needed for the court certificate (BSA 2023 s.63). e-Sakshya / CCTNS link: planned.'),
    ('Offline (design)', ' - no network? Test saved on phone with its hash, sent later; sent twice = saved once; '
                         'phone lost = admin blocks the login (planned).'),
], size=11.5, head=15)

p.save(OUT)
print('saved', OUT)
