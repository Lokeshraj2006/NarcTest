"""v19: built on v18. Slide 5: the certificate picture and its text are replaced by one short line with a link
to a full sample report on Google Drive (shared as 'anyone with the link'); the map goes back to full width."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SRC = '../deck_v18/K26234_SIH26231_v18.pptx'
OUT = 'K26234_SIH26231_v19.pptx'
REPORT = 'https://drive.google.com/file/d/1NIR2OyYZRdmEM1R3cz9ghb7rZqyY1D1Z/view?usp=sharing'
HEAD, INK, LINK = RGBColor(0x79, 0x2E, 0x00), RGBColor(0, 0, 0), RGBColor(0x05, 0x63, 0xC1)
E = 914400
p = Presentation(SRC)
s5 = p.slides[4]


def sid(i):
    return next(s for s in s5.shapes if s.shape_id == i)


def drop(shape):
    shape._element.getparent().remove(shape._element)


# remove the certificate block added in v18
drop(next(s for s in s5.shapes if s.name == 'Section 63 page'))
for s in list(s5.shapes):
    if s.has_text_frame and s.text_frame.text.strip().startswith(('Court Certificate', 'Done:')):
        drop(s)

# map back to full width, a little less sky and road so the link line fits under it
fut, mp = sid(11), sid(19)
L, T, R, B = 0.178, 0.14, 0.015, 0.43
src = mp._element.find('.//' + qn('a:blipFill')).find(qn('a:srcRect'))
for k, v in zip('ltrb', (L, T, R, B)):
    src.set(k, str(int(v * 100000)))
mp.left, mp.width = fut.left, fut.width
mp.height = int(fut.width * (530 * (1 - T - B)) / (960 * (1 - L - R)))

# one line under the map: court certificate + link to the sample report
x, y, w, h = fut.left, mp.top + mp.height + Inches(0.05), fut.width, Inches(0.44)
tb = s5.shapes.add_textbox(x, y, w, h)
tf = tb.text_frame
tf.word_wrap, tf.auto_size = True, MSO_AUTO_SIZE.NONE
tf.margin_left = tf.margin_right = Inches(0.03)
tf.margin_top = tf.margin_bottom = Inches(0)
par = tf.paragraphs[0]
for text, kind in (('Court Certificate: ', 'head'),
                   ('every PDF report now has a filled BSA section 63 page (photo hash, phone, time, GPS) '
                    'for the officer and expert to sign. ', None),
                   ('Open a full sample report (PDF, 2 pages)', 'link')):
    r = par.add_run()
    r.text = text
    r.font.name, r.font.size, r.font.bold = 'Calibri', Pt(13), kind == 'head'
    r.font.color.rgb = HEAD if kind == 'head' else (LINK if kind == 'link' else INK)
    if kind == 'link':
        r.font.underline = True
        r.hyperlink.address = REPORT

# PowerPoint's PDF export can drop text links, so an invisible click area covers the line too
box = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
box.name = 'Sample report link area'
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
sf = box.fill._xPr.find(qn('a:solidFill'))[0]
sf.append(sf.makeelement(qn('a:alpha'), {'val': '0'}))
box.line.fill.background()
box.click_action.hyperlink.address = REPORT

p.save(OUT)
print('saved', OUT, '| map h', round(mp.height / E, 2), 'line y', round(y / E, 2))
