"""v14: built on v13. NCRB figure cited on the slide; Future Scope turned into a dated plan
(Tamil is now done in the app, 29 Sep 2026)."""
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN

SRC = '../deck_v13/K26234_SIH26231_v13.pptx'
OUT = 'K26234_SIH26231_v14.pptx'
p = Presentation(SRC)


def sid(slide, i):
    return next(s for s in slide.shapes if s.shape_id == i)


def set_para(shape, prefix, text):
    for par in shape.text_frame.paragraphs:
        if par.text.strip().startswith(prefix) and par.runs:
            par.runs[0].text = text
            for r in par.runs[1:]:
                r.text = ''
            return
    raise SystemExit('not found: ' + prefix)


s2, s5 = p.slides[1], p.slides[4]

# Slide 2 - cite the NCRB number on the slide
sid(s2, 41).text_frame.paragraphs[0].runs[0].text = \
    'India: 1,20,010 NDPS cases in 2023 (source: NCRB, Crime in India 2023)'

# Slide 5 - Future Scope as a plan with months
fs = sid(s5, 22)
set_para(fs, 'Offline Mode', 'Offline Mode (month 1) - save the test on phone, send it when network is back.')
set_para(fs, 'Phone Security', "Phone Security (month 2) - sign each test with the phone's security chip.")
set_para(fs, 'Police System Integration', 'Police System Integration (after pilot) - apply to NIC or MHA to link with '
                                          'e-Sakshya and CCTNS. Our CSV export is ready.')
set_para(fs, 'Lab Comparison', 'Lab Comparison (during pilot) - compare field results with FSL lab results.')
set_para(fs, 'Regional Languages', 'Regional Languages - Tamil is done in the app. Hindi and voice help in month 3.')
head = sid(s5, 28)
head.text_frame.paragraphs[0].runs[0].text = 'Future Scope (our plan after SIH)'
for r in head.text_frame.paragraphs[0].runs[1:]:
    r.text = ''
head.text_frame.word_wrap = False
head.text_frame.auto_size = MSO_AUTO_SIZE.NONE
head.width = Inches(5.5)
head.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT

p.save(OUT)
print('saved', OUT)
