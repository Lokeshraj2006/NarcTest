"""v7: declutter after the 60/90 review ('deck is dense, small text'). Built on v6.

Slide 3  tighter line spacing in Measured Results -> bigger Proof of Results photos
Slide 5  Privacy + Chain of Custody blocks moved to SPEAKER NOTES (for Q&A); QR screenshots back to full size;
         map a bit smaller; tamper test redrawn in 2 rows with bigger text
"""
from pptx import Presentation
from pptx.util import Inches, Pt

SRC = '../deck_v6/K26234_SIH26231_v6.pptx'
V5 = '../deck_v5/K26234_SIH26231_v5.pptx'
OUT = 'K26234_SIH26231_v7.pptx'
TALL = '../ppt_content/evidence/tamper_proof_tall.png'
p = Presentation(SRC)
E = 914400


def sid(slide, i):
    return next(s for s in slide.shapes if s.shape_id == i)


def starts(slide, text):
    return next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip().startswith(text))


def place(shape, x=None, y=None, w=None, h=None):
    if x is not None: shape.left = Inches(x)
    if y is not None: shape.top = Inches(y)
    if w is not None: shape.width = Inches(w)
    if h is not None: shape.height = Inches(h)


def drop(shape):
    shape._element.getparent().remove(shape._element)


s3, s5 = p.slides[2], p.slides[4]
v5s5 = Presentation(V5).slides[4]

# ------------------------------------------------------------------ Slide 3: bigger proof photos
for par in sid(s3, 19).text_frame.paragraphs:   # Measured Results bullets: same size, tighter lines
    par.line_spacing = 1.1
place(sid(s3, 20), y=5.42)                        # simulated-swatch note
place(sid(s3, 23), y=6.36)                        # divider
card, ref_head = sid(s3, 13), sid(s3, 22)
proof_head = starts(s3, 'Proof of Results')
proof = next(s for s in s3.shapes if s.name == 'Proof grid')
top, hh = 6.84, 3.50
pw = hh * proof.width / proof.height
place(proof, x=16.55 - pw, y=top, w=pw, h=hh)
place(proof_head, x=16.55 - pw, y=6.42, w=pw)
cw = 16.55 - pw - 0.15 - 10.15
ch = cw * card.height / card.width
place(card, x=10.15, y=top + (hh - ch) / 2, w=cw, h=ch)
place(ref_head, x=10.15, y=6.42, w=cw)

# ------------------------------------------------------------------ Slide 5: blocks -> speaker notes
priv = starts(s5, 'Privacy & How Long')
chain = starts(s5, 'Chain of Custody')
notes = 'Q&A NOTES (not on slide)\n\n' + priv.text_frame.text + '\n\n' + chain.text_frame.text
drop(priv)
drop(chain)
s5.notes_slide.notes_text_frame.text = notes
for i in (10, 15, 16, 30):                        # QR screenshots + heading back to their v5 size/place
    a, b = sid(s5, i), sid(v5s5, i)
    a.left, a.top, a.width, a.height = b.left, b.top, b.width, b.height

fut, mp = sid(s5, 11), sid(s5, 19)
fx, fw = fut.left / E, fut.width / E
mh = 2.0
mw = mh * mp.width / mp.height
dx = Inches(fx + (fw - mw) / 2) - mp.left
place(mp, x=fx + (fw - mw) / 2, y=2.52, w=mw, h=mh)
for i in (29, 31):                                # "Map - Geo-Tagged test result" follows the map
    sid(s5, i).left += dx
old = next(s for s in s5.shapes if s.name == 'Tamper proof')
drop(old)
place(starts(s5, 'Tamper Test'), y=4.62)
place(starts(s5, 'real test on our saved'), y=4.66)
pic = s5.shapes.add_picture(TALL, Inches(fx), Inches(5.02), width=Inches(fw))
pic.name = 'Tamper proof'
bottom = 5.02 + pic.height / E
place(fut, y=bottom + 0.12, h=10.25 - bottom - 0.12)
place(sid(s5, 28), y=bottom + 0.18)               # Future Scope heading
place(sid(s5, 22), y=bottom + 0.58)               # Future Scope text
for par in sid(s5, 22).text_frame.paragraphs:
    par.line_spacing = 1.12
    for r in par.runs:
        r.font.size = Pt(16)

p.save(OUT)
print('saved', OUT)
