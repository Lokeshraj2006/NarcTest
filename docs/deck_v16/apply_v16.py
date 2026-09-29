"""v16: built on v15. Answers the latest 60/90 review with 3 small changes:
Slide 3  real-scans image now also shows the colour gap (Delta E) numbers from the 19 scans
Slide 4  pilot cost broken down; one grey note on legal status (presumptive) and privacy
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE

SRC = '../deck_v15/K26234_SIH26231_v15.pptx'
OUT = 'K26234_SIH26231_v16.pptx'
SCANS = '../ppt_content/evidence/real_scans.png'
GREY = RGBColor(0x55, 0x55, 0x55)
E = 914400
p = Presentation(SRC)


def sid(slide, i):
    return next(s for s in slide.shapes if s.shape_id == i)


s3, s4 = p.slides[2], p.slides[3]

# Slide 3 - new real-scans image (same width, a bit taller), kept above the footer
old = next(s for s in s3.shapes if s.name == 'Real scans')
x, y = old.left, old.top
old._element.getparent().remove(old._element)
w = Inches(6.55)
pic = s3.shapes.add_picture(SCANS, x, y, width=w)
pic.name = 'Real scans'
bottom = Inches(10.36)
if pic.top + pic.height > bottom:                     # shrink to fit, keep it centred in the column
    h = bottom - pic.top
    nw = int(pic.width * h / pic.height)
    pic.left = x + (w - nw) // 2
    pic.width, pic.height = nw, h

# Slide 4 - cost breakdown in the pilot plan
pilot = sid(s4, 55)
cost = next(par for par in pilot.text_frame.paragraphs if par.text.startswith('Cost'))
cost.runs[1].text = (' - 20 cards at ₹10 = ₹200, server ₹3,000 a month for 3 months = ₹9,000. '
                     'Total about ₹9,200 (estimate)')
for r in cost.runs[2:]:
    r.text = ''

# Slide 4 - legal status + privacy, as a grey side note next to "Feasibility & Sustainability"
head = sid(s4, 20)
nx = head.left / E + head.width / E + 0.15
tb = s4.shapes.add_textbox(Inches(nx), Inches(head.top / E + 0.02), Inches(9.55 - nx), Inches(0.46))
tf = tb.text_frame
tf.word_wrap, tf.auto_size = True, MSO_AUTO_SIZE.NONE
tf.margin_left = tf.margin_right = Inches(0.03)
tf.margin_top = tf.margin_bottom = Inches(0.01)
for i, (label, text) in enumerate([('Legal: ', 'the app gives a first check only, the FSL lab confirms it.'),
                                   ('Privacy: ', 'no face or suspect name is stored, only a sample code.')]):
    par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    for t, b in ((label, True), (text, False)):
        r = par.add_run()
        r.text = t
        r.font.name, r.font.size, r.font.bold = 'Calibri', Pt(12), b
        r.font.color.rgb = GREY

p.save(OUT)
print('saved', OUT, 'scan pic', round(pic.left / E, 2), round(pic.top / E, 2), round(pic.width / E, 2), round(pic.height / E, 2))
